"""CSV reports built from completed workflow runs.

Shared by campaign-, workflow-, and organization-usage-scoped reports.
The DB client supplies the row set; this module owns the column layout
so every endpoint emits the same shape.
"""

import csv
import io
from datetime import UTC, datetime
from typing import Any, List, Optional

from api.db import db_client
from api.utils.artifacts import artifact_url


def _collect_extracted_variable_keys(runs: List[Any]) -> list[str]:
    """Collect all unique extracted variable keys across runs, preserving insertion order."""
    keys: dict[str, None] = {}
    for run in runs:
        gathered = run.gathered_context or {}
        extracted = gathered.get("extracted_variables", {})
        if isinstance(extracted, dict):
            for key in extracted:
                keys.setdefault(key, None)
    return list(keys)


def build_run_report_csv(runs: List[Any]) -> io.StringIO:
    """Build a CSV from completed workflow runs."""
    extracted_var_keys = _collect_extracted_variable_keys(runs)

    output = io.StringIO()
    writer = csv.writer(output)

    pre_headers = [
        "Run ID",
        "Campaign ID",
        "Agent ID",
        "Agent Definition ID",
        "Created At",
        "Phone Number",
        "Call Disposition",
        "Call Duration (s)",
    ]
    post_headers = [
        "Call Tags",
        "Transcript URL",
        "Recording URL",
    ]
    writer.writerow(pre_headers + extracted_var_keys + post_headers)

    for run in runs:
        initial = run.initial_context or {}
        gathered = run.gathered_context or {}
        usage = run.usage_info or {}

        call_tags = gathered.get("call_tags", [])
        if isinstance(call_tags, list):
            call_tags = ", ".join(str(t) for t in call_tags)

        pre_values = [
            run.id,
            run.campaign_id if run.campaign_id is not None else "",
            run.workflow_id,
            run.definition_id if run.definition_id is not None else "",
            run.created_at.isoformat() if run.created_at else "",
            initial.get("phone_number", ""),
            gathered.get("mapped_call_disposition", ""),
            usage.get("call_duration_seconds", ""),
        ]

        extracted = gathered.get("extracted_variables", {})
        if not isinstance(extracted, dict):
            extracted = {}
        extracted_values = [extracted.get(key, "") for key in extracted_var_keys]

        post_values = [
            call_tags,
            artifact_url(run.public_access_token, "transcript") or "",
            artifact_url(run.public_access_token, "recording") or "",
        ]

        writer.writerow(pre_values + extracted_values + post_values)

    output.seek(0)
    return output


async def generate_campaign_report_csv(
    campaign_id: int,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
) -> tuple[io.StringIO, str]:
    """Generate a CSV report for a campaign."""
    runs = await db_client.get_completed_runs_for_report(
        campaign_id=campaign_id, start_date=start_date, end_date=end_date
    )
    return build_run_report_csv(runs), f"campaign_{campaign_id}_report.csv"


async def generate_workflow_report_csv(
    workflow_id: int,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
) -> tuple[io.StringIO, str]:
    """Generate a CSV report for all completed runs of a workflow."""
    runs = await db_client.get_completed_runs_for_report(
        workflow_id=workflow_id, start_date=start_date, end_date=end_date
    )
    return build_run_report_csv(runs), f"workflow_{workflow_id}_report.csv"


async def generate_usage_runs_report_csv(
    organization_id: int,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    filters: Optional[list[dict]] = None,
) -> tuple[io.StringIO, str]:
    """Generate a CSV report for runs visible on the org-wide usage page.

    Honors the same date / filter inputs as the `/usage/runs` listing.
    """
    runs = await db_client.get_usage_runs_for_report(
        organization_id,
        start_date=start_date,
        end_date=end_date,
        filters=filters,
    )
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    return build_run_report_csv(runs), f"usage_runs_{timestamp}.csv"


async def generate_usage_runs_report_excel(
    organization_id: int,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    filters: Optional[list[dict]] = None,
) -> tuple[bytes, str]:
    """Generate an Excel report for runs visible on the org-wide usage page.

    Honors the same date / filter inputs as the `/usage/runs` listing.
    """
    import openpyxl
    from io import BytesIO
    from api.utils.transcript import detect_user_intent, detect_user_intent_async, generate_transcript_text
    from api.utils.artifacts import artifact_url
    from openpyxl.styles import Alignment, Font

    runs = await db_client.get_usage_runs_for_report(
        organization_id,
        start_date=start_date,
        end_date=end_date,
        filters=filters,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Call Reports"

    # Headers
    headers = [
        "Name",
        "Number",
        "Call Type",
        "Disposition",
        "Intent",
        "Status",
        "Duration (seconds)",
        "Call Tags",
        "Recording URL",
        "Recording Transcript URL",
        "Full Recording Text"
    ]
    ws.append(headers)

    ws.column_dimensions['A'].width = 20  # Name
    ws.column_dimensions['B'].width = 15  # Number
    ws.column_dimensions['C'].width = 15  # Call Type
    ws.column_dimensions['D'].width = 18  # Disposition
    ws.column_dimensions['E'].width = 16  # Intent
    ws.column_dimensions['F'].width = 15  # Status
    ws.column_dimensions['G'].width = 18  # Duration
    ws.column_dimensions['H'].width = 20  # Call Tags
    ws.column_dimensions['I'].width = 40  # Rec URL
    ws.column_dimensions['J'].width = 40  # Transcript URL
    ws.column_dimensions['K'].width = 60  # Full Recording Text

    header_font = Font(name="Calibri", size=11, bold=True)
    for col_num in range(1, 12):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    row_idx = 1
    for run in runs:
        row_idx += 1
        # 1. Name
        initial = run.initial_context or {}
        name = initial.get("name") or initial.get("customer_name") or initial.get("first_name", "")
        if initial.get("last_name"):
            name = f"{name} {initial.get('last_name')}".strip()
        if not name:
            name = "N/A"

        # 2. Number & Call Type
        call_type = run.call_type or initial.get("direction") or "outbound"
        caller_number = initial.get("caller_number")
        called_number = initial.get("called_number") or initial.get("phone_number")
        if call_type == "inbound":
            number = caller_number
        else:
            number = initial.get("phone_number") or called_number

        # 3. Disposition
        disposition = None
        if run.gathered_context:
            disposition = run.gathered_context.get("mapped_call_disposition")
        if not disposition and run.logs:
            logs_field = run.logs
            callbacks = []
            if isinstance(logs_field, dict):
                callbacks = logs_field.get("telephony_status_callbacks", [])
            elif isinstance(logs_field, list):
                for item in logs_field:
                    if isinstance(item, dict) and "telephony_status_callbacks" in item:
                        callbacks = item.get("telephony_status_callbacks", [])
                        break
            if callbacks and isinstance(callbacks, list):
                disposition = callbacks[-1].get("status")
        if not disposition:
            disposition = run.state or ""

        # 4. Duration
        duration_str = run.usage_info.get("call_duration_seconds", "0") if run.usage_info else "0"
        try:
            duration = float(duration_str)
        except (ValueError, TypeError):
            duration = 0.0

        # Determine status
        disposition_lower = disposition.lower()
        is_received = False
        if duration > 0:
            is_received = True
        elif disposition_lower in (
            "completed",
            "answered",
            "user_hangup",
            "agent_hangup",
            "call_duration_exceeded",
            "user_qualified",
            "transfer_call",
        ):
            is_received = True
        
        status = "Received" if is_received else "Not Received"

        # 5. Tags
        tags = ""
        if run.gathered_context:
            call_tags = run.gathered_context.get("call_tags", [])
            if isinstance(call_tags, list):
                tags = ", ".join(str(t) for t in call_tags)

        # 6 & 7. URLs
        token = run.public_access_token
        rec_url = artifact_url(token, "recording") if token else ""
        transcript_url = artifact_url(token, "transcript") if token else ""

        # 8. Full Recording Text
        logs = run.logs or {}
        if isinstance(logs, dict):
            events = logs.get("realtime_feedback_events") or []
            transcript_text = generate_transcript_text(events)
        elif isinstance(logs, list):
            transcript_text = generate_transcript_text(logs)
        else:
            transcript_text = ""

        # 9. Intent Detection
        intent = detect_user_intent(
            gathered_context=run.gathered_context,
            transcript_text=transcript_text,
            disposition=disposition,
            duration=duration,
        )

        ws.append([
            name,
            number,
            call_type,
            disposition,
            intent,
            status,
            duration,
            tags,
            rec_url,
            transcript_url,
            transcript_text
        ])

        # Apply alignment (wrap text for full recording text)
        ws.cell(row=row_idx, column=11).alignment = Alignment(wrap_text=True, vertical="top")

    bytes_io = BytesIO()
    wb.save(bytes_io)
    bytes_io.seek(0)

    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    return bytes_io.getvalue(), f"usage_runs_{timestamp}.xlsx"
