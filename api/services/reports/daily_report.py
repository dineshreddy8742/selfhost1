from datetime import datetime, time
from typing import Any, Dict, List, Optional
from zoneinfo import ZoneInfo

from api.db import db_client


class DailyReportService:
    async def get_daily_report(
        self,
        organization_id: int,
        date: str,
        timezone: str,
        workflow_id: Optional[int] = None,
    ) -> Dict[str, Any]:
        """
        Get daily report for a specific date and timezone.

        Args:
            organization_id: The organization ID to filter by
            date: Date in YYYY-MM-DD format
            timezone: IANA timezone string (e.g., "America/New_York")
            workflow_id: Optional workflow ID to filter by (None means all workflows)
        """
        # Parse date and timezone
        tz = ZoneInfo(timezone)
        date_obj = datetime.strptime(date, "%Y-%m-%d")

        # Create start and end datetime in the specified timezone
        start_dt = datetime.combine(date_obj, time.min, tzinfo=tz)
        end_dt = datetime.combine(date_obj, time.max, tzinfo=tz)

        # Convert to UTC for database queries
        start_utc = start_dt.astimezone(ZoneInfo("UTC"))
        end_utc = end_dt.astimezone(ZoneInfo("UTC"))

        # Get workflow runs from database (optimized - only required fields)
        runs = await db_client.get_workflow_runs_for_daily_report(
            organization_id=organization_id,
            start_utc=start_utc,
            end_utc=end_utc,
            workflow_id=workflow_id,
        )

        # Calculate metrics
        total_runs = len(runs)
        xfer_count = sum(
            1
            for run in runs
            if run["gathered_context"]
            and run["gathered_context"].get("mapped_call_disposition") == "XFER"
        )

        # Calculate disposition distribution
        disposition_counts = {}
        for run in runs:
            if run["gathered_context"]:
                disposition = run["gathered_context"].get(
                    "mapped_call_disposition", "UNKNOWN"
                )
                disposition_counts[disposition] = (
                    disposition_counts.get(disposition, 0) + 1
                )

        # Sort dispositions by count and get top 5
        sorted_dispositions = sorted(
            disposition_counts.items(), key=lambda x: x[1], reverse=True
        )

        disposition_distribution = []
        other_count = 0

        for i, (disposition, count) in enumerate(sorted_dispositions):
            if i < 5:
                disposition_distribution.append(
                    {
                        "disposition": disposition,
                        "count": count,
                        "percentage": round(
                            (count / total_runs * 100) if total_runs > 0 else 0, 2
                        ),
                    }
                )
            else:
                other_count += count

        # Add "Other" category if there are more than 5 dispositions
        if other_count > 0:
            disposition_distribution.append(
                {
                    "disposition": "Other",
                    "count": other_count,
                    "percentage": round(
                        (other_count / total_runs * 100) if total_runs > 0 else 0, 2
                    ),
                }
            )

        # Calculate call duration distribution
        duration_buckets = {
            "0-10": {"range_start": 0, "range_end": 10, "count": 0},
            "10-30": {"range_start": 10, "range_end": 30, "count": 0},
            "30-60": {"range_start": 30, "range_end": 60, "count": 0},
            "60-120": {"range_start": 60, "range_end": 120, "count": 0},
            "120-180": {"range_start": 120, "range_end": 180, "count": 0},
            ">180": {"range_start": 180, "range_end": None, "count": 0},
        }

        for run in runs:
            if run["usage_info"]:
                duration_str = run["usage_info"].get("call_duration_seconds")
                if duration_str:
                    try:
                        duration = float(duration_str)
                        if duration < 10:
                            duration_buckets["0-10"]["count"] += 1
                        elif duration < 30:
                            duration_buckets["10-30"]["count"] += 1
                        elif duration < 60:
                            duration_buckets["30-60"]["count"] += 1
                        elif duration < 120:
                            duration_buckets["60-120"]["count"] += 1
                        elif duration < 180:
                            duration_buckets["120-180"]["count"] += 1
                        else:
                            duration_buckets[">180"]["count"] += 1
                    except (ValueError, TypeError):
                        pass

        # Format duration distribution
        call_duration_distribution = []
        total_calls_with_duration = sum(b["count"] for b in duration_buckets.values())

        for bucket_name, bucket_data in duration_buckets.items():
            call_duration_distribution.append(
                {
                    "bucket": bucket_name,
                    "range_start": bucket_data["range_start"],
                    "range_end": bucket_data["range_end"],
                    "count": bucket_data["count"],
                    "percentage": round(
                        (bucket_data["count"] / total_calls_with_duration * 100)
                        if total_calls_with_duration > 0
                        else 0,
                        2,
                    ),
                }
            )

        return {
            "date": date,
            "timezone": timezone,
            "workflow_id": workflow_id,
            "metrics": {"total_runs": total_runs, "xfer_count": xfer_count},
            "disposition_distribution": disposition_distribution,
            "call_duration_distribution": call_duration_distribution,
        }

    async def get_workflows_for_organization(
        self, organization_id: int
    ) -> List[Dict[str, Any]]:
        """
        Get all workflows for an organization.
        """
        workflows = await db_client.get_workflows_for_organization(organization_id)

        return [{"id": workflow.id, "name": workflow.name} for workflow in workflows]

    async def get_daily_runs_detail(
        self,
        organization_id: int,
        date: str,
        timezone: str,
        workflow_id: Optional[int] = None,
    ) -> List[Dict[str, Any]]:
        """
        Get detailed workflow runs for CSV export.

        Args:
            organization_id: The organization ID to filter by
            date: Date in YYYY-MM-DD format
            timezone: IANA timezone string (e.g., "America/New_York")
            workflow_id: Optional workflow ID to filter by
        """
        # Parse date and timezone
        tz = ZoneInfo(timezone)
        date_obj = datetime.strptime(date, "%Y-%m-%d")

        # Create start and end datetime in the specified timezone
        start_dt = datetime.combine(date_obj, time.min, tzinfo=tz)
        end_dt = datetime.combine(date_obj, time.max, tzinfo=tz)

        # Convert to UTC for database queries
        start_utc = start_dt.astimezone(ZoneInfo("UTC"))
        end_utc = end_dt.astimezone(ZoneInfo("UTC"))

        # Get workflow runs from database (optimized - only required fields)
        runs = await db_client.get_workflow_runs_for_daily_report(
            organization_id=organization_id,
            start_utc=start_utc,
            end_utc=end_utc,
            workflow_id=workflow_id,
        )

        # Format runs for CSV export
        detailed_runs = []
        for run in runs:
            # Phone number is already extracted at the database level
            # Try customer_phone_number first, then fall back to initial_context
            phone_number = run["gathered_context"].get(
                "customer_phone_number", ""
            ) or run["initial_context"].get("phone_number", "")

            # Disposition is already extracted at the database level
            disposition = run["gathered_context"].get("mapped_call_disposition", "")

            # Duration is already extracted at the database level
            duration_seconds = 0
            duration_str = run["usage_info"].get("call_duration_seconds", "0")
            try:
                duration_seconds = float(duration_str)
            except (ValueError, TypeError):
                duration_seconds = 0

            detailed_runs.append(
                {
                    "phone_number": phone_number,
                    "disposition": disposition,
                    "duration_seconds": duration_seconds,
                    "workflow_id": run["workflow_id"],
                    "run_id": run["id"],
                    "workflow_name": run["workflow_name"],
                    "created_at": run["created_at"].isoformat(),
                }
            )

        return detailed_runs

    async def get_daily_runs_excel(
        self,
        organization_id: int,
        date: str,
        timezone: str,
        workflow_id: Optional[int] = None,
    ) -> bytes:
        """
        Get daily runs and export them as an Excel workbook (.xlsx).
        """
        import openpyxl
        from io import BytesIO
        from api.utils.artifacts import artifact_url
        from api.utils.transcript import detect_user_intent, detect_user_intent_async, generate_transcript_text

        # Parse date and timezone
        tz = ZoneInfo(timezone)
        date_obj = datetime.strptime(date, "%Y-%m-%d")

        # Create start and end datetime in the specified timezone
        start_dt = datetime.combine(date_obj, time.min, tzinfo=tz)
        end_dt = datetime.combine(date_obj, time.max, tzinfo=tz)

        # Convert to UTC for database queries
        start_utc = start_dt.astimezone(ZoneInfo("UTC"))
        end_utc = end_dt.astimezone(ZoneInfo("UTC"))

        # Fetch the detailed runs with full fields (logs, access token etc)
        runs = await db_client.get_workflow_runs_for_excel_report(
            organization_id=organization_id,
            start_utc=start_utc,
            end_utc=end_utc,
            workflow_id=workflow_id,
        )

        # Create virtual Excel workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Call Reports"

        # Headers
        headers = [
            "Name",
            "Number",
            "Call Type",
            "Disposition",
            "Intent",
            "Status",
            "Duration (seconds)",
            "Call Tags",
            "Recording URL",
            "Recording Transcript URL",
            "Full Recording Text"
        ]
        ws.append(headers)

        # Format column widths and wrap text for transcript
        ws.column_dimensions['A'].width = 20  # Name
        ws.column_dimensions['B'].width = 15  # Number
        ws.column_dimensions['C'].width = 15  # Call Type
        ws.column_dimensions['D'].width = 18  # Disposition
        ws.column_dimensions['E'].width = 16  # Intent
        ws.column_dimensions['F'].width = 15  # Status
        ws.column_dimensions['G'].width = 18  # Duration
        ws.column_dimensions['H'].width = 20  # Call Tags
        ws.column_dimensions['I'].width = 40  # Rec URL
        ws.column_dimensions['J'].width = 40  # Transcript URL
        ws.column_dimensions['K'].width = 60  # Full Recording Text

        from openpyxl.styles import Alignment, Font

        # Header style
        header_font = Font(name="Calibri", size=11, bold=True)
        for col_num in range(1, 12):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Add data rows
        for run in runs:
            # 1. Name
            initial = run.get("initial_context") or {}
            name = initial.get("name") or initial.get("customer_name") or initial.get("first_name", "")
            if initial.get("last_name"):
                name = f"{name} {initial.get('last_name')}".strip()
            if not name:
                name = "N/A"

            # 2. Number & Call Type
            call_type = run.get("call_type") or initial.get("direction") or "outbound"
            caller_number = initial.get("caller_number")
            called_number = initial.get("called_number") or initial.get("phone_number")
            if call_type == "inbound":
                number = caller_number
            else:
                number = initial.get("phone_number") or called_number

            # 3. Disposition
            disposition = None
            if run.get("gathered_context"):
                disposition = run.get("gathered_context", {}).get("mapped_call_disposition")
            if not disposition and run.get("logs"):
                logs_field = run.get("logs")
                callbacks = []
                if isinstance(logs_field, dict):
                    callbacks = logs_field.get("telephony_status_callbacks", [])
                elif isinstance(logs_field, list):
                    for item in logs_field:
                        if isinstance(item, dict) and "telephony_status_callbacks" in item:
                            callbacks = item.get("telephony_status_callbacks", [])
                            break
                if callbacks and isinstance(callbacks, list):
                    disposition = callbacks[-1].get("status")
            if not disposition:
                disposition = run.get("state") or ""

            # 4. Duration
            duration_str = run.get("usage_info", {}).get("call_duration_seconds", "0")
            try:
                duration = float(duration_str)
            except (ValueError, TypeError):
                duration = 0.0

            # 5. Status (Received / Not Received)
            disposition_lower = disposition.lower()
            is_received = False
            if duration > 0:
                is_received = True
            elif disposition_lower in (
                "completed",
                "answered",
                "user_hangup",
                "agent_hangup",
                "call_duration_exceeded",
                "user_qualified",
                "transfer_call",
            ):
                is_received = True
            status = "Received" if is_received else "Not Received"

            # 6. Tags
            tags = ""
            if run.get("gathered_context"):
                call_tags = run.get("gathered_context", {}).get("call_tags", [])
                if isinstance(call_tags, list):
                    tags = ", ".join(str(t) for t in call_tags)

            # 7 & 8. URLs
            token = run.get("public_access_token")
            rec_url = artifact_url(token, "recording") or ""
            transcript_url = artifact_url(token, "transcript") or ""

            # 9. Full Recording Text
            logs = run.get("logs") or {}
            if isinstance(logs, dict):
                events = logs.get("realtime_feedback_events") or []
                transcript_text = generate_transcript_text(events)
            elif isinstance(logs, list):
                transcript_text = generate_transcript_text(logs)
            else:
                transcript_text = ""

            # 10. Intent Detection
            intent = await detect_user_intent_async(
                gathered_context=run.get("gathered_context"),
                transcript_text=transcript_text,
                disposition=disposition,
                duration=duration,
                organization_id=organization_id,
            )

            ws.append([
                name,
                number,
                call_type,
                disposition,
                intent,
                status,
                duration,
                tags,
                rec_url,
                transcript_url,
                transcript_text
            ])

            # Apply alignment (wrap text for full recording text)
            row_idx = ws.max_row
            ws.cell(row=row_idx, column=11).alignment = Alignment(wrap_text=True, vertical="top")

        # Write workbook to bytes
        bytes_io = BytesIO()
        wb.save(bytes_io)
        bytes_io.seek(0)
        return bytes_io.getvalue()

